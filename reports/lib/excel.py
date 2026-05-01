"""Excel workbook builder.

Per docs/04 § 10: one sheet per report; "Methodology" sheet at the front
documenting period, filters, n, and known caveats.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)


def new_workbook() -> Workbook:
    wb = Workbook()
    # remove the auto sheet; reports add explicit sheets
    wb.remove(wb.active)
    return wb


def add_methodology_sheet(
    wb: Workbook,
    *,
    report_title: str,
    period: str,
    filters: list[str],
    population: list[str],
    caveats: list[str],
    freshness: dict,
) -> None:
    ws = wb.create_sheet("Methodology", 0)
    rows = [
        (report_title, ""),
        ("", ""),
        ("Generated at (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")),
        ("Data last synced (UTC)", freshness.get("last_sync_utc") or "—"),
        ("Data age (hours)", str(freshness.get("age_hours") or "—")),
        ("Period", period),
        ("", ""),
        ("Filters applied", ""),
    ]
    for f in filters:
        rows.append(("", f))
    rows.append(("", ""))
    rows.append(("Population", ""))
    for p in population:
        rows.append(("", p))
    rows.append(("", ""))
    rows.append(("Caveats / known data issues", ""))
    for c in caveats:
        rows.append(("", c))

    for r_idx, (k, v) in enumerate(rows, start=1):
        ws.cell(row=r_idx, column=1, value=k)
        ws.cell(row=r_idx, column=2, value=v)
        if r_idx == 1:
            ws.cell(row=r_idx, column=1).font = TITLE_FONT

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80


def add_table_sheet(
    wb: Workbook,
    *,
    name: str,
    headers: list[str],
    rows: list[list],
) -> None:
    ws = wb.create_sheet(name)
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    # auto-width-ish based on header
    for c_idx, h in enumerate(headers, start=1):
        col = get_column_letter(c_idx)
        ws.column_dimensions[col].width = max(12, len(str(h)) + 4)


def save(wb: Workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
